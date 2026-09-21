"""
Escala Blitz - DC-PI2
Aplicação Flask para visualizar e editar a escala de funcionários
da Blitz (terceirizada que presta serviço para a iMile Delivery).

Como rodar:
    pip install -r requirements.txt
    python app.py
Depois acesse http://localhost:5000

Os dados ficam salvos em data/escala_data.json (arquivo simples,
sem necessidade de banco de dados). Basta fazer backup desse
arquivo de vez em quando.
"""
import json
import os
import calendar
from datetime import date

from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, send_file, flash, session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from io import BytesIO

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_PATH = os.path.join(APP_DIR, "data", "escala_data.json")

app = Flask(__name__)
app.secret_key = os.environ.get("ESCALA_SECRET_KEY", "escala-blitz-dc-pi2")  # troque se for expor publicamente

# Senha única de acesso ao site. Pode trocar aqui direto ou, melhor, definir
# a variável de ambiente ESCALA_SENHA antes de rodar (assim não fica escrita
# no código). Ex.: no Windows -> set ESCALA_SENHA=minhasenha
#                   no Linux/Mac -> export ESCALA_SENHA=minhasenha
LOGIN_PASSWORD = os.environ.get("ESCALA_SENHA", "blitz2026")


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("logado"):
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        senha = request.form.get("senha", "")
        if senha == LOGIN_PASSWORD:
            session["logado"] = True
            destino = request.form.get("next") or url_for("index")
            return redirect(destino)
        flash("Senha incorreta.")
    return render_template("login.html", next=request.args.get("next", ""))


@app.route("/logout")
def logout():
    session.pop("logado", None)
    return redirect(url_for("login"))

STATUS_OPTIONS = ["", "T", "DSR", "Folga", "Falta", "Atestado", "Abono"]
STATUS_LABELS = {"": "—", "T": "T", "DSR": "DSR", "Folga": "Folga",
                  "Falta": "Falta", "Atestado": "Atestado", "Abono": "Abono"}
WEEKDAY_SHORT = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]
MONTH_NAMES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
               "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]

STATUS_COLORS_XLSX = {
    "T": "E4F4EA", "DSR": "E7EEF6", "Folga": "EFECF6", "Falta": "FBE8E6",
    "Atestado": "FBEED9", "Abono": "F6F1D8", "": "ECEFF2",
}


# ---------------------------------------------------------------------------
# Persistência (arquivo JSON simples)
# ---------------------------------------------------------------------------
def load_data():
    with open(DATA_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def save_data(data):
    with open(DATA_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def days_in_month_keys(month_key):
    y, m = (int(x) for x in month_key.split("-"))
    n = calendar.monthrange(y, m)[1]
    return [f"{month_key}-{d:02d}" for d in range(1, n + 1)]


def weekday_short(date_key):
    y, m, d = (int(x) for x in date_key.split("-"))
    return WEEKDAY_SHORT[date(y, m, d).weekday()]


def ensure_days(op, month_key):
    days = days_in_month_keys(month_key)
    for emp in op["employees"]:
        for dk in days:
            emp["dias"].setdefault(dk, "")
    return days


def get_month(data, month_key):
    for m in data["months"]:
        if m["key"] == month_key:
            return m
    return None


def compute_summary(op, days):
    rows = []
    totals = {"Folga": 0, "Falta": 0, "Atestado": 0, "Abono": 0, "DSR": 0, "T": 0}
    for emp in op["employees"]:
        c = {"Folga": 0, "Falta": 0, "Atestado": 0, "Abono": 0, "DSR": 0, "T": 0}
        for dk in days:
            v = emp["dias"].get(dk, "")
            if v in c:
                c[v] += 1
        total = sum(c.values())
        status = "OK" if total == len(days) else "NOT"
        for k in totals:
            totals[k] += c[k]
        rows.append({"nome": emp["nome"], "c": c, "total": total, "status": status})
    return {"rows": rows, "totals": totals, "total_days": len(days)}


# ---------------------------------------------------------------------------
# Rotas
# ---------------------------------------------------------------------------
@app.route("/")
@login_required
def index():
    data = load_data()
    if not data["months"]:
        return render_template("index.html", data=data, month=None)
    first_key = data["months"][0]["key"]
    return redirect(url_for("ver_mes", month_key=first_key, op_idx=0))


@app.route("/mes/<month_key>/op/<int:op_idx>")
@login_required
def ver_mes(month_key, op_idx):
    data = load_data()
    month = get_month(data, month_key)
    if month is None:
        flash("Mês não encontrado.")
        return redirect(url_for("index"))
    op_idx = max(0, min(op_idx, len(month["operations"]) - 1))
    op = month["operations"][op_idx]
    days = ensure_days(op, month_key)
    day_labels = [(dk, weekday_short(dk), dk[-2:]) for dk in days]
    summary = compute_summary(op, days)
    return render_template(
        "index.html", data=data, month=month, month_key=month_key,
        op=op, op_idx=op_idx, days=days, day_labels=day_labels,
        summary=summary, status_options=STATUS_OPTIONS, status_labels=STATUS_LABELS,
    )


@app.route("/mes/<month_key>/op/<int:op_idx>/salvar", methods=["POST"])
@login_required
def salvar(month_key, op_idx):
    data = load_data()
    month = get_month(data, month_key)
    op = month["operations"][op_idx]
    days = ensure_days(op, month_key)

    for i, emp in enumerate(op["employees"]):
        emp["nome"] = request.form.get(f"nome_{i}", emp["nome"]).strip() or emp["nome"]
        emp["turno"] = request.form.get(f"turno_{i}", emp["turno"]).strip()
        emp["cargo"] = request.form.get(f"cargo_{i}", emp["cargo"]).strip()
        for dk in days:
            val = request.form.get(f"dia_{i}_{dk}", "")
            if val in STATUS_OPTIONS:
                emp["dias"][dk] = val

    save_data(data)
    flash("Alterações salvas.")
    return redirect(url_for("ver_mes", month_key=month_key, op_idx=op_idx))


@app.route("/mes/<month_key>/op/<int:op_idx>/funcionario/adicionar", methods=["POST"])
@login_required
def adicionar_funcionario(month_key, op_idx):
    data = load_data()
    month = get_month(data, month_key)
    op = month["operations"][op_idx]
    days = days_in_month_keys(month_key)
    op["employees"].append({
        "nome": "NOVO FUNCIONÁRIO", "turno": "", "cargo": "",
        "dias": {dk: "" for dk in days},
    })
    save_data(data)
    return redirect(url_for("ver_mes", month_key=month_key, op_idx=op_idx))


@app.route("/mes/<month_key>/op/<int:op_idx>/funcionario/<int:emp_idx>/remover", methods=["POST"])
@login_required
def remover_funcionario(month_key, op_idx, emp_idx):
    data = load_data()
    month = get_month(data, month_key)
    op = month["operations"][op_idx]
    if 0 <= emp_idx < len(op["employees"]):
        op["employees"].pop(emp_idx)
        save_data(data)
    return redirect(url_for("ver_mes", month_key=month_key, op_idx=op_idx))


@app.route("/mes/adicionar", methods=["POST"])
@login_required
def adicionar_mes():
    data = load_data()
    month_key = request.form.get("month_key")  # formato AAAA-MM
    copiar = request.form.get("copiar") == "sim"
    if not month_key or get_month(data, month_key):
        flash("Informe um mês válido que ainda não exista.")
        return redirect(url_for("index"))

    y, m = (int(x) for x in month_key.split("-"))
    label = f"{MONTH_NAMES[m - 1]} {y}"
    days = days_in_month_keys(month_key)

    employees = []
    if copiar and data["months"]:
        src_op = data["months"][-1]["operations"][0]
        for e in src_op["employees"]:
            employees.append({
                "nome": e["nome"], "turno": e["turno"], "cargo": e["cargo"],
                "dias": {dk: "" for dk in days},
            })

    data["months"].append({"key": month_key, "label": label,
                            "operations": [{"name": "Geral", "employees": employees}]})
    data["months"].sort(key=lambda m: m["key"])
    save_data(data)
    return redirect(url_for("ver_mes", month_key=month_key, op_idx=0))


@app.route("/mes/<month_key>/operacao/adicionar", methods=["POST"])
@login_required
def adicionar_operacao(month_key):
    data = load_data()
    month = get_month(data, month_key)
    nome = request.form.get("nome", "").strip()
    if not nome:
        flash("Informe o nome da operação.")
        return redirect(url_for("ver_mes", month_key=month_key, op_idx=0))
    month["operations"].append({"name": nome, "employees": []})
    save_data(data)
    return redirect(url_for("ver_mes", month_key=month_key, op_idx=len(month["operations"]) - 1))


@app.route("/mes/<month_key>/op/<int:op_idx>/exportar")
@login_required
def exportar(month_key, op_idx):
    data = load_data()
    month = get_month(data, month_key)
    op = month["operations"][op_idx]
    days = ensure_days(op, month_key)

    wb = Workbook()
    ws = wb.active
    ws.title = (month["label"] + " " + op["name"])[:31]

    header = ["Nome", "Turno", "Cargo"] + days
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for emp in op["employees"]:
        row = [emp["nome"], emp["turno"], emp["cargo"]] + [emp["dias"].get(dk, "") for dk in days]
        ws.append(row)

    # colore as células de status
    for r in range(2, ws.max_row + 1):
        for c in range(4, 4 + len(days)):
            cell = ws.cell(row=r, column=c)
            color = STATUS_COLORS_XLSX.get(cell.value, "FFFFFF")
            cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    ws.append([])
    summary = compute_summary(op, days)
    ws.append(["Funcionário", "Folgas", "Faltas", "Atestados", "Abono", "DSR", "Dias Trab.", "Total Dias", "Status"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for r in summary["rows"]:
        ws.append([r["nome"], r["c"]["Folga"], r["c"]["Falta"], r["c"]["Atestado"],
                   r["c"]["Abono"], r["c"]["DSR"], r["c"]["T"], r["total"], r["status"]])

    for col in ws.columns:
        length = max((len(str(cell.value)) for cell in col if cell.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(length + 2, 6), 20)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Escala_Blitz_{month_key}_{op['name'].replace(' ', '_')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename,
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(debug=True, port=5000)
